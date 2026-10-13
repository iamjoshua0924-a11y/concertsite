#!/usr/bin/env python3
"""기존 Apps Script 사이트의 구글 시트(xlsx로 내려받은 것)를 새 사이트 가져오기 JSON으로 변환한다.

사용법:
    pip install openpyxl
    python3 scripts/convert_apps_script.py 시트.xlsx import-data.json [--images 이미지폴더] [--event-date 2027-01-30]

- 이미지폴더에는 드라이브 파일 ID를 파일명으로 둔다 (예: 1grIhtx....png). 없는 이미지는 건너뛴다.
- 접속 토큰 값은 옮기지 않는다. 라벨·역할·팀만 옮기고, 새 사이트에서 새 코드를 발급한다.
- 결과 JSON에는 연락처·운영규약이 들어 있으므로 저장소에 커밋하지 말 것(.gitignore 처리됨).
"""
import argparse
import base64
import json
import mimetypes
import re
import sys
from pathlib import Path

import openpyxl

# 테마 색상이 비어 있는 팀에 순서대로 배정 (팀 대표가 나중에 바꿀 수 있음)
FALLBACK_COLORS = ['#1D4FB5', '#2BA88A', '#E0A21B', '#8A4FD8', '#E0567A', '#3A9BD9']
STATUS_MAP = {'완료': 'done', '작성중': 'doing', '진행중': 'doing'}
VENUE_LINKS = [
    {'label': '감마홀 대관규정', 'url': 'https://docs.google.com/document/d/1Cq2PCD17weWWKD5TmxWx_DLBlT7i5TzeU187Q6gUfNQ/edit'},
    {'label': '감마홀 장비 리스트', 'url': 'https://docs.google.com/spreadsheets/d/1O7KXUGc4Ucg1iJlQjDxFCNFCa-ZJFfDEYBnOIvkD_zQ/edit'},
]


def rows(wb, name):
    if name not in wb.sheetnames:
        return []
    it = wb[name].iter_rows(values_only=True)
    header = [str(h) if h is not None else '' for h in next(it, [])]
    out = []
    for r in it:
        if not any(c not in (None, '') for c in r):
            continue
        out.append({h: ('' if v is None else v) for h, v in zip(header, r)})
    return out


def s(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip() if v is not None else ''


def visible(row):
    return s(row.get('isVisible') or 'Y') != 'N'


def lines(v):
    return [x.strip() for x in s(v).split('\n') if x.strip()]


def shrink(path, max_side=1800):
    """Pillow가 있으면 큰 이미지를 긴 변 1800px JPEG로 줄인다(모바일 로딩용)."""
    raw = path.read_bytes()
    try:
        from io import BytesIO
        from PIL import Image
    except ImportError:
        return raw, None
    if len(raw) < 350 * 1024:
        return raw, None
    im = Image.open(BytesIO(raw))
    if im.mode in ('RGBA', 'LA', 'P'):
        bg = Image.new('RGB', im.size, 'white')
        im = im.convert('RGBA')
        bg.paste(im, mask=im.split()[-1])
        im = bg
    im.thumbnail((max_side, max_side))
    buf = BytesIO()
    im.convert('RGB').save(buf, 'JPEG', quality=85, optimize=True)
    out = buf.getvalue()
    return (out, 'image/jpeg') if len(out) < len(raw) else (raw, None)


def load_image(images_dir, file_id, name):
    if not images_dir:
        return None
    for p in Path(images_dir).glob(f'{file_id}.*'):
        mime = mimetypes.guess_type(p.name)[0] or 'image/jpeg'
        data, new_mime = shrink(p)
        if new_mime:
            mime = new_mime
            name = re.sub(r'\.\w+$', '', name or p.stem) + '.jpg'
        return {'name': name or p.name, 'mime': mime, 'base64': base64.b64encode(data).decode()}
    print(f'  이미지 없음, 건너뜀: {file_id} ({name})', file=sys.stderr)
    return None


def parse_times(text):
    """'오후 6시~9시(팀당 40분 내외 배정)' → ('18:00', '21:00', 40)"""
    start = end = None
    slot = None
    m = re.search(r'(오전|오후)?\s*(\d{1,2})\s*시\s*[~\-–]\s*(\d{1,2})\s*시', text)
    if m:
        pm = m.group(1) == '오후'
        a, b = int(m.group(2)), int(m.group(3))
        if pm and a < 12:
            a += 12
        if pm and b < 12:
            b += 12
        start, end = f'{a:02d}:00', f'{b:02d}:00'
    m = re.search(r'팀당\s*(\d{1,3})\s*분', text)
    if m:
        slot = int(m.group(1))
    return start, end, slot


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx')
    ap.add_argument('out')
    ap.add_argument('--images')
    ap.add_argument('--event-date', default='')
    ap.add_argument('--title', default='')
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    ev = (rows(wb, 'EventInfo') or [{}])[0]
    start, end, slot = parse_times(s(ev.get('eventTime')))
    schedules = lines(ev.get('schedules'))

    teams = []
    images_by_team = {}
    for img in sorted(rows(wb, 'TeamImages'), key=lambda r: float(r.get('sortOrder') or 0)):
        if visible(img):
            images_by_team.setdefault(s(img['teamId']), []).append(img)
    setlists = {}
    for row in sorted(rows(wb, 'Setlists'), key=lambda r: float(r.get('orderNo') or 0)):
        setlists.setdefault(s(row['teamId']), []).append({
            'title': s(row.get('songTitle')), 'artist': s(row.get('artist')), 'duration': s(row.get('runningTime')),
            'cue': s(row.get('cueSheet')), 'note': s(row.get('note')),
        })
    color_i = 0
    for t in sorted(rows(wb, 'Teams'), key=lambda r: float(r.get('sortOrder') or 0)):
        tid = s(t['teamId'])
        color = s(t.get('themeColor'))
        if not re.fullmatch(r'#[0-9a-fA-F]{6}', color):
            color = FALLBACK_COLORS[color_i % len(FALLBACK_COLORS)]
            color_i += 1
        imgs = [load_image(a.images, s(i['fileId']), s(i['fileName'])) for i in images_by_team.get(tid, [])]
        teams.append({
            'id': re.sub(r'[^a-z0-9-]', '-', tid.lower()), 'name': s(t.get('teamName')) or tid,
            'intro': s(t.get('teamIntro')), 'snsUrl': s(t.get('snsUrl')), 'memberCount': s(t.get('memberCount')) or None,
            'instruments': s(t.get('instruments')), 'leaderName': s(t.get('leaderName')),
            'leaderContact': s(t.get('leaderContact')), 'themeColor': color,
            'setlist': setlists.get(tid, []), 'images': [i for i in imgs if i],
        })
    id_map = {s(t['teamId']): re.sub(r'[^a-z0-9-]', '-', s(t['teamId']).lower()) for t in rows(wb, 'Teams')}
    if teams and teams[0]['memberCount'] is not None:
        for t in teams:
            t['memberCount'] = int(float(t['memberCount'])) if t['memberCount'] else None

    posters = [load_image(a.images, s(p['fileId']), s(p['fileName'])) for p in rows(wb, 'EventPosters') if visible(p)]

    progress_by_col = {}
    for p in rows(wb, 'CollectionProgress'):
        progress_by_col.setdefault(s(p['collectionId']), {})[id_map.get(s(p['teamId']), s(p['teamId']))] = {
            'status': STATUS_MAP.get(s(p.get('status')), 'todo'), 'note': s(p.get('note')),
            'updatedBy': s(p.get('updatedBy')), 'updatedAt': s(p.get('updatedAt')),
        }
    collections = [{
        'title': s(c['title']), 'description': s(c.get('description')), 'dueDate': s(c.get('dueDate'))[:10],
        'targetTeamIds': [id_map.get(x, x) for x in lines(c.get('targetTeamIds'))],
        'status': 'closed' if s(c.get('status')) == '마감' else 'open',
        'progress': progress_by_col.get(s(c['collectionId']), {}), 'createdBy': s(c.get('createdBy')), 'createdAt': s(c.get('createdAt')),
    } for c in rows(wb, 'Collections') if visible(c)]

    options = {}
    for o in sorted(rows(wb, 'PollOptions'), key=lambda r: float(r.get('sortOrder') or 0)):
        options.setdefault(s(o['pollId']), []).append(s(o['optionLabel']))
    import datetime
    today = datetime.date.today().isoformat()

    def poll_due(p):
        due = s(p.get('dueDate'))[:10]
        # 기존 사이트는 마감일을 강제하지 않아 '진행중' 투표가 마감일을 넘긴 채 쓰이고 있음 → 마감일만 비움
        return '' if (due and due < today and s(p.get('status')) != '마감') else due

    polls = [{
        'title': s(p['title']), 'description': s(p.get('description')),
        'targetTeamId': '' if s(p.get('targetTeamId')) in ('', 'all') else id_map.get(s(p['targetTeamId']), ''),
        'options': options.get(s(p['pollId']), []), 'anonymous': s(p.get('allowAnonymousVote')) == 'Y',
        'dueDate': poll_due(p), 'status': 'closed' if s(p.get('status')) == '마감' else 'open',
        'createdBy': s(p.get('createdBy')), 'createdAt': s(p.get('createdAt')),
    } for p in rows(wb, 'Polls') if visible(p)]

    data = {
        'format': 'concertsite-import@1',
        'event': {
            'title': a.title or s(ev.get('title')),
            'venue': s(ev.get('venue')),
            'eventDate': a.event_date,
            'startTime': start or '18:00', 'endTime': end or '21:00', 'slotMinutes': slot or 40, 'changeoverMinutes': 10,
            'talkRoomUrl': s(ev.get('talkRoomUrl')), 'notice': s(ev.get('notice')),
            'scheduleNote': '\n'.join(schedules), 'links': VENUE_LINKS,
        },
        'posters': [p for p in posters if p],
        'rules': s((rows(wb, 'Rules') or [{}])[0].get('content')),
        'teams': teams,
        'collections': collections,
        'polls': polls,
        'meetings': [{'title': s(m['title']), 'content': s(m.get('content')), 'author': s(m.get('author')),
                      'pinned': s(m.get('isPinned')) == 'Y', 'createdAt': s(m.get('createdAt'))} for m in rows(wb, 'MeetingNotes')],
        'questions': [{'question': s(q['question']), 'answer': s(q.get('answer')), 'author': s(q.get('author')),
                       'createdAt': s(q.get('createdAt'))} for q in rows(wb, 'Questions')],
        'posts': [{'title': s(p['title']), 'content': s(p.get('content')), 'author': s(p.get('author')),
                   'passwordHash': s(p.get('passwordHash')), 'createdAt': s(p.get('createdAt'))} for p in rows(wb, 'FreePosts') if visible(p)],
        # 토큰 값(tokenValue)은 의도적으로 제외
        'tokens': [{'label': s(t.get('tokenLabel')), 'role': s(t.get('role')), 'teamId': id_map.get(s(t.get('teamId')), ''), 'note': s(t.get('note'))}
                   for t in rows(wb, 'AccessTokens') if s(t.get('isActive') or 'Y') != 'N'],
    }
    Path(a.out).write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding='utf-8')
    print(f'팀 {len(teams)}, 셋리스트 {sum(len(t["setlist"]) for t in teams)}곡, 이미지 {sum(len(t["images"]) for t in teams)}+포스터 {len(data["posters"])}, '
          f'취합 {len(collections)}, 투표 {len(polls)}, 게시글 {len(data["posts"])}, 접속코드 {len(data["tokens"])}개 → {a.out}')


if __name__ == '__main__':
    main()
